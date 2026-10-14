"""Update the v2 deck in place-style: text edits keep run formatting; drops 2 slides (12 -> 10)."""
import re
import sys
from pptx import Presentation

SRC = r'D:\Sem 7\FOAI\FOAI case study\Clinical_Trial_MAS_Case_Study_v2.pptx'
DST = r'D:\Sem 7\FOAI\FOAI case study\Clinical_Trial_MAS_Case_Study_v4.pptx'

prs = Presentation(SRC)
slides = list(prs.slides)


def set_para(p, text):
    runs = p.runs
    if not runs:
        return
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def find(slide, startswith):
    for sh in slide.shapes:
        if sh.has_text_frame and sh.text_frame.text.startswith(startswith):
            return sh
    raise KeyError(startswith)


def set_shape_lines(sh, lines):
    paras = sh.text_frame.paragraphs
    assert len(paras) == len(lines), (sh.text_frame.text, len(paras), len(lines))
    for p, t in zip(paras, lines):
        set_para(p, t)


def replace(slide, old_start, new):
    sh = find(slide, old_start)
    lines = new if isinstance(new, list) else [new]
    set_shape_lines(sh, lines)


def set_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


s1, s2, s3, s4, s5, s6, s7, s8, s9, s10, s11, s12 = slides

# 1 title
replace(s1, 'Presented by', 'Presented by: Druva Teja Korlakunta (CB.SC.U4CSE23317)  |  Manvitha Dungi (CB.SC.U4CSE23318)')
set_notes(s1, 'One-line pitch: four agents with private constraints must all approve a single trial protocol. '
              'Introduction: real protocol design is a negotiation; a single model gives one averaged answer, needs every '
              "party's confidential limits in one place, and nobody can veto. Our MAS keeps four independent agents, "
              'each seeing only its own constraints, with an IRB veto and an independent Constraint Checker.')

# 3 PEAS
replace(s3, 'Consensus reached', ['Consensus reached (all four APPROVE)',
                                 'Violations = 0; power >= 0.80; cost <= budget',
                                 'Rounds, proposals accepted vs blocked, objections'])
replace(s3, 'Shared protocol document', ['Shared protocol document + public board',
                                       'Public calculators: power, cost, recruitment',
                                       'Guarded private state for each agent'])
replace(s3, 'APPROVE / REJECT', ['APPROVE / REJECT / COUNTER_PROPOSE',
                                'Amendments to own fields only (output -> field)',
                                'Qualitative requests and objections (no numbers)'])
replace(s3, 'Current protocol version', ['Current protocol version and public history',
                                        'Public tool report for any draft (input -> view)',
                                        'Feedback on own blocked proposals; own private limits'])
set_notes(s3, 'System-level PEAS. Performance is measured by the independent Constraint Checker, not by the agents. '
              'Metrics: consensus = all four approve the same version; violations = hard-constraint failures on the final '
              'protocol (target 0); power >= 0.80; cost <= budget. Input -> sensor: protocol, history, tool report, '
              'feedback, own private state. Output -> actuator: Sponsor edits sample/hospitals/arms/budget/remote '
              'monitoring; Biostatistician edits sample/randomisation/endpoint/alpha; IRB edits dosage/monitoring/AE '
              'reporting/exclusions; Advocate edits inclusion/visits/procedures.')

# 4 environment
tbl = [sh for sh in s4.shapes if sh.has_table][0].table
rows = {
    5: ('Dynamics', 'Static between rounds', 'Changes only via accepted amendments; interim scenario shifts limits'),
    6: ('State space', 'Discrete, known rules', 'Finite protocol fields and public formulas; peers\' limits hidden'),
}
for r, vals in rows.items():
    for c, v in enumerate(vals):
        set_para(tbl.cell(r, c).text_frame.paragraphs[0], v)
set_notes(s4, 'Classification follows the Russell and Norvig task-environment dimensions. Agent category: goal-based '
              'deliberative agents with a model of the shared protocol; the Sponsor adds a learning component. System '
              'type: mixed cooperative-competitive MAS. Mock mode is deterministic; no external dataset, all parameters '
              'are synthetic.')

# 5 agents
set_notes(s5, 'Agent category: goal-based, utility-oriented deliberative agents. The Sponsor is additionally a learning '
              'agent: it infers the hidden power bar from public objections only. IRB safety amendments are binding; '
              'the Patient Advocate is advisory but nothing is final without its approval.')

# 9 tools / structure
replace(s9, 'Python 3.11', 'Python 3.10+ standard library; any PC, no GPU')
replace(s9, 'pytest, 18 tests', 'pytest: 96 tests (unit, algorithm, integration)')
replace(s9, 'Dataclasses; new scenario', 'Protocol + PrivateState dataclasses; synthetic scenarios')
set_notes(s9, 'About 3,200 lines of Python including tests across 12 modules. Scenarios are derived from the baseline by '
              'a small modify() function, so a new scenario is a few lines; a new agent is one subclass + prompt file + '
              'PrivateState + constraint check. Input representation: Protocol dataclass (sample size, hospitals, dosage, '
              'arms, duration, budget, power, alpha, randomisation, endpoint, visits, remote monitoring, procedures, '
              'inclusion/exclusion, monitoring, AE reporting). No dataset: all values are synthetic. CLI flags: --mode, '
              '--scenario, --llm, --max-rounds, --summary.')

# 11 results
replace(s11, 'All eight reach', 'All eight reach CONSENSUS within 10 rounds. 96 automated tests pass.')
set_notes(s11, 'From python main.py --mode compare --scenario all (mock mode), re-run and verified. Test levels: unit '
               '(tools, constraints, models), algorithm (authority order, partial acceptance, deadlock, max rounds), '
               'integration (private-state guard, no prompt leakage, LLM validation/fallback), scenario (all eight reach '
               'valid consensus). 96 passed. The single-agent baseline is intentionally naive; its violation count '
               'illustrates silent trade-offs without vetoes, not a claim that no single optimiser could satisfy the '
               'constraints.')

# 12 closing
set_notes(s12, 'Likely questions: Why not one LLM? (privacy, veto, audit). Is the single-agent comparison fair? (it is '
               'intentionally naive; the point is architectural). What if no agreement exists? (DEADLOCK or '
               'NO_CONSENSUS with unresolved conflicts listed). Limitations: synthetic models, hand-written mock '
               'strategies, fixed processing order. Demo script: python main.py; --scenario high_risk_drug --summary; '
               '--mode compare --scenario all; pytest -q.')


def drop(prs, slide):
    lst = prs.slides._sldIdLst
    for sldId in list(lst):
        if prs.part.related_part(sldId.rId) is slide.part:
            prs.part.drop_rel(sldId.rId)
            lst.remove(sldId)
            return
    raise RuntimeError('slide not found')


# 12 -> 10 slides: overview (2) is covered by title + agent slide; architecture (6) by layered module slide + workflow
drop(prs, s2)
drop(prs, s6)

# renumber kicker labels "NN  //  ..." to match the new order
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_text_frame and re.match(r'^\d\d  //  ', sh.text_frame.text):
            p = sh.text_frame.paragraphs[0]
            for r in p.runs:
                if re.match(r'^\d\d', r.text):
                    r.text = re.sub(r'^\d\d', f'{i:02d}', r.text, count=1)
                    break

prs.save(DST)
print('saved', DST, len(prs.slides))
